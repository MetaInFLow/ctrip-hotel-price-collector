#!/usr/bin/env python3
"""Install and validate runtime components bundled with the native runner."""

from __future__ import annotations

import os
import platform
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Callable, Mapping
from urllib.parse import urlparse

from ctrip_runtime import bundled_browser_binary, cloakbrowser_cache_dir


PROXY_ENVIRONMENT_VARIABLES = (
    "HTTPS_PROXY",
    "https_proxy",
    "HTTP_PROXY",
    "http_proxy",
    "ALL_PROXY",
    "all_proxy",
)


def invalid_proxy_environment_variable(
    environ: Mapping[str, str] | None = None,
) -> str | None:
    """Return a malformed proxy variable name without exposing its value."""

    environment = os.environ if environ is None else environ
    for name in PROXY_ENVIRONMENT_VARIABLES:
        value = environment.get(name)
        if not value:
            continue
        try:
            parsed = urlparse(value)
            _ = parsed.port
        except ValueError:
            return name
        if parsed.scheme not in {"http", "https", "socks5", "socks5h"}:
            return name
        if not parsed.hostname:
            return name

    try:
        import httpx

        client = httpx.Client()
        client.close()
    except Exception:
        if environment.get("no_proxy"):
            return "no_proxy"
        if environment.get("NO_PROXY"):
            return "NO_PROXY"
        return "代理配置"
    return None


def _probe_browser_binary(
    browser_binary: Path,
    *,
    run_command: Callable[..., Any] = subprocess.run,
) -> None:
    command = [str(browser_binary), "--version"]
    if platform.system() == "Windows":
        command.append("--no-startup-window")
    result = run_command(
        command,
        capture_output=True,
        text=True,
        timeout=15,
        check=False,
    )
    if result.returncode != 0:
        detail = str(result.stderr or result.stdout or "").strip()
        raise RuntimeError(f"CloakBrowser 无法启动：{detail or 'unknown error'}")


def _probe_playwright_driver(
    *,
    playwright_factory: Callable[[], Any] | None = None,
) -> None:
    if playwright_factory is None:
        from playwright.sync_api import sync_playwright

        playwright_factory = sync_playwright

    manager = playwright_factory()
    playwright = manager.start()
    playwright.stop()


def _probe_excel_engine() -> None:
    from openpyxl import Workbook, load_workbook

    with tempfile.TemporaryDirectory() as temporary_directory:
        workbook_path = Path(temporary_directory) / "runtime-check.xlsx"
        workbook = Workbook()
        workbook.active.append(["运行环境", "OK"])
        workbook.save(workbook_path)
        loaded = load_workbook(workbook_path, read_only=True)
        try:
            if loaded.active["B1"].value != "OK":
                raise RuntimeError("Excel 回读结果不一致")
        finally:
            loaded.close()


def setup_runtime(
    *,
    cache_dir: Path | None = None,
    ensure_binary_fn: Callable[[], str] | None = None,
    run_command: Callable[..., Any] = subprocess.run,
    playwright_factory: Callable[[], Any] | None = None,
) -> Path:
    """Locate the browser and prove the bundled native runner is usable."""

    target_cache = Path(cache_dir or cloakbrowser_cache_dir()).expanduser().resolve()
    target_cache.mkdir(parents=True, exist_ok=True)
    os.environ["CLOAKBROWSER_CACHE_DIR"] = str(target_cache)
    os.environ.setdefault("CLOAKBROWSER_AUTO_UPDATE", "false")

    packaged_browser = bundled_browser_binary()
    if packaged_browser is not None:
        os.environ.setdefault("CLOAKBROWSER_BINARY_PATH", str(packaged_browser))

    local_browser_configured = bool(os.environ.get("CLOAKBROWSER_BINARY_PATH"))
    if not local_browser_configured:
        invalid_proxy = invalid_proxy_environment_variable()
        if invalid_proxy:
            raise RuntimeError(
                f"代理环境变量 {invalid_proxy} 的地址格式无效。"
                "请修正系统代理，或清除该变量后重新运行安装。"
            )

    try:
        import cloakbrowser  # noqa: F401
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise RuntimeError(f"运行程序缺少内置组件：{exc.name}") from exc

    try:
        _probe_playwright_driver(playwright_factory=playwright_factory)
    except Exception as exc:
        raise RuntimeError(f"内置浏览器驱动无法启动：{exc}") from exc
    try:
        _probe_excel_engine()
    except Exception as exc:
        raise RuntimeError(f"内置 Excel 引擎无法读写：{exc}") from exc

    if ensure_binary_fn is None:
        from cloakbrowser.download import ensure_binary

        ensure_binary_fn = ensure_binary

    try:
        browser_binary = Path(ensure_binary_fn()).expanduser().resolve()
    except Exception as exc:
        if local_browser_configured:
            raise RuntimeError(
                f"CloakBrowser 本地浏览器组件无法加载（{type(exc).__name__}）。"
                "请重新解压完整客户包后重试。"
            ) from None
        raise RuntimeError(
            f"CloakBrowser 浏览器组件下载失败（{type(exc).__name__}）。"
            "请确认网络可访问 CloakBrowser 下载服务后重试。"
        ) from None
    if not browser_binary.is_file():
        raise RuntimeError(f"CloakBrowser 浏览器文件不存在：{browser_binary}")

    _probe_browser_binary(browser_binary, run_command=run_command)
    return browser_binary
