#!/usr/bin/env python3
"""把携程酒店采集结果导出为 Excel（openpyxl 实现）。

本脚本是唯一的 Excel 生成入口，使用纯 Python 的 openpyxl 生成四个工作表，
无需 Node.js 或其它运行时。

用法：
    python ctrip_hotel_excel_builder.py \
        --input-dir output/ctrip_hotel_prices \
        --output output/ctrip_hotel_prices/ctrip_hotel_prices.xlsx
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRICE_HEADERS = [
    "酒店名称", "入住日期", "离店日期", "房型", "房型ID", "销售方案", "销售方案ID",
    "价格", "接口价格", "价格来源", "原价", "货币", "显示价格", "总价展示", "床型", "早餐及权益", "取消政策",
    "预订状态", "可预订", "余房", "销售方案Key", "接口状态", "接口URL", "JSON文件", "采集时间", "详情页URL",
    "页面价格文本", "页面序号", "页面价格XPath", "页面匹配方式", "展示所有房型XPath",
]
SUMMARY_HEADERS = [
    "酒店名称", "入住日期", "离店日期", "状态", "价格模式", "接口响应数", "房型行数",
    "页面抽查状态", "页面抽查数", "价格差异数", "详情页URL", "JSON文件", "错误",
]
RESPONSE_HEADERS = ["酒店名称", "入住日期", "离店日期", "接口状态", "请求方法", "接口URL", "JSON文件", "响应JSON字符数"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")

DATE_HEADERS = {"入住日期", "离店日期"}
CURRENCY_HEADERS = {"价格", "接口价格", "原价"}
INTEGER_HEADERS = {"余房", "销售方案ID", "接口状态", "响应JSON字符数", "页面序号", "接口响应数", "房型行数", "页面抽查数", "价格差异数"}


def _excel_date(value):
    """把 yyyy-mm-dd 字符串转为 datetime.date，供 openpyxl 落成日期单元格。"""
    if not value:
        return None
    s = str(value).strip()
    if len(s) == 10:
        try:
            return datetime.strptime(s, "%Y-%m-%d").date()
        except ValueError:
            return value
    return value


def _cell_value(value, header):
    if header in DATE_HEADERS:
        return _excel_date(value)
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _style_sheet(ws, headers, rows, widths, wrap_cols=()):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    for row in rows:
        ws.append([_cell_value(v, h) for h, v in zip(headers, row)])
    ws.freeze_panes = "A2"

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for col_name in wrap_cols:
        if col_name in headers:
            idx = headers.index(col_name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def _apply_number_formats(ws, headers, row_count):
    if not row_count:
        return
    for header in DATE_HEADERS:
        if header in headers:
            idx = headers.index(header) + 1
            for r in range(2, row_count + 2):
                ws.cell(row=r, column=idx).number_format = "yyyy-mm-dd"
    for header in CURRENCY_HEADERS:
        if header in headers:
            idx = headers.index(header) + 1
            for r in range(2, row_count + 2):
                ws.cell(row=r, column=idx).number_format = "#,##0.00"
    for header in INTEGER_HEADERS:
        if header in headers:
            idx = headers.index(header) + 1
            for r in range(2, row_count + 2):
                ws.cell(row=r, column=idx).number_format = "#,##0"


def _walk_json_files(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*.json") if p.is_file())


def _read_collection(input_dir: Path):
    files = _walk_json_files(input_dir)
    index = {"items": []}
    records = []
    response_rows = []

    for file in files:
        payload = json.loads(file.read_text(encoding="utf-8"))
        if file.name == "index.json":
            index = payload
            continue
        if file.name.endswith(".error.json"):
            continue
        for row in payload.get("room_rows", []):
            record = dict(row)
            record["JSON文件"] = str(file.relative_to(input_dir))
            records.append(record)
        for response in payload.get("responses", []):
            response_rows.append({
                "酒店名称": payload.get("hotel_name", ""),
                "入住日期": payload.get("check_in", ""),
                "离店日期": payload.get("check_out", ""),
                "接口状态": response.get("status"),
                "请求方法": response.get("method", ""),
                "接口URL": response.get("url", ""),
                "JSON文件": str(file.relative_to(input_dir)),
                "响应JSON字符数": len(json.dumps(response.get("data", {}), ensure_ascii=False)),
            })

    return index, records, response_rows


def build(input_dir: Path, output_path: Path) -> None:
    index, records, response_rows = _read_collection(input_dir)

    price_rows = [[record.get(h) for h in PRICE_HEADERS] for record in records]

    summary_rows = []
    for item in index.get("items", []):
        summary_rows.append([
            item.get("hotel_name", ""),
            item.get("check_in", ""),
            item.get("check_out", ""),
            item.get("status", ""),
            item.get("price_mode", "response"),
            item.get("response_count", 0),
            item.get("room_row_count", 0),
            item.get("page_price_check_status", "skipped"),
            item.get("page_price_check_count", 0),
            item.get("page_price_mismatch_count", 0),
            item.get("detail_url", ""),
            Path(item.get("file", "")).name if item.get("file") else "",
            item.get("error", ""),
        ])

    response_values = [[r.get(h) for h in RESPONSE_HEADERS] for r in response_rows]

    wb = Workbook()

    ws_price = wb.active
    ws_price.title = "房型价格"
    _style_sheet(
        ws_price, PRICE_HEADERS, price_rows,
        [24, 12, 12, 24, 14, 20, 14, 12, 12, 12, 12, 10, 14, 14, 20, 28, 28, 12, 10, 10, 24, 12, 38, 48, 24, 52, 24, 12, 48, 16, 48],
        wrap_cols=["早餐及权益", "取消政策", "接口URL", "JSON文件", "详情页URL", "页面价格文本", "页面价格XPath", "展示所有房型XPath"],
    )
    _apply_number_formats(ws_price, PRICE_HEADERS, len(price_rows))

    ws_summary = wb.create_sheet("采集汇总")
    _style_sheet(
        ws_summary, SUMMARY_HEADERS, summary_rows,
        [24, 12, 12, 12, 16, 14, 12, 18, 12, 12, 52, 52, 28],
        wrap_cols=["详情页URL", "JSON文件", "错误"],
    )
    _apply_number_formats(ws_summary, SUMMARY_HEADERS, len(summary_rows))

    ws_resp = wb.create_sheet("接口概览")
    _style_sheet(
        ws_resp, RESPONSE_HEADERS, response_values,
        [24, 12, 12, 12, 12, 52, 52, 18],
        wrap_cols=["接口URL", "JSON文件"],
    )
    _apply_number_formats(ws_resp, RESPONSE_HEADERS, len(response_values))

    ws_note = wb.create_sheet("说明")
    ws_note.append(["项目", "内容"])
    for cell in ws_note[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    now = datetime.now()
    note_rows = [
        ("生成时间", now),
        ("原始 JSON 目录", str(input_dir)),
        ("数据说明", "房型价格可来自 getHotelRoomListInland 接口响应或页面 XPath；每行的价格来源列标明实际口径，完整接口响应和页面抽查记录仍保存在 JSON 文件中。"),
        ("价格口径", "response 模式的价格列来自接口 priceInfo.price；page_xpath 模式的价格列来自配置的页面价格 XPath，接口价格列保留接口原值。"),
        ("登录状态", "采集器只在检测到“我的订单”且登录入口消失后继续。"),
    ]
    for label, value in note_rows:
        ws_note.append([label, value])
    ws_note.cell(row=2, column=2).number_format = "yyyy-mm-dd hh:mm"
    ws_note.column_dimensions["A"].width = 22
    ws_note.column_dimensions["B"].width = 90
    for r in range(2, ws_note.max_row + 1):
        ws_note.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Excel 已生成：{output_path}")
    print(f"房型明细行：{len(price_rows)}；采集记录：{len(summary_rows)}；接口响应：{len(response_values)}")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="把携程酒店采集结果导出为 Excel（纯 Python openpyxl 实现）。",
    )
    parser.add_argument("--input-dir", required=True, help="采集结果 JSON 目录（output/ctrip_hotel_prices）")
    parser.add_argument("--output", required=True, help="输出 .xlsx 路径")
    args = parser.parse_args(argv)

    build(
        Path(args.input_dir).expanduser().resolve(),
        Path(args.output).expanduser().resolve(),
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
