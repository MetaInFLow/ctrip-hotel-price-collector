import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BUILDER_PATH = PROJECT_ROOT / "scripts" / "ctrip_hotel_excel_builder.py"


def load_builder_module():
    spec = importlib.util.spec_from_file_location("ctrip_hotel_excel_builder", BUILDER_PATH)
    if spec is None or spec.loader is None:
        raise AssertionError(f"无法加载 Excel 生成脚本：{BUILDER_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class ExcelExportTests(unittest.TestCase):
    def test_python_builder_writes_readable_workbook(self):
        builder = load_builder_module()

        with tempfile.TemporaryDirectory() as temporary_directory:
            input_dir = Path(temporary_directory) / "input"
            input_dir.mkdir()
            (input_dir / "index.json").write_text(
                json.dumps(
                    {
                        "items": [
                            {
                                "hotel_name": "测试酒店",
                                "check_in": "2026-09-03",
                                "check_out": "2026-09-04",
                                "status": "ok",
                                "response_count": 1,
                                "room_row_count": 1,
                                "detail_url": "https://hotels.ctrip.com/hotels/1.html?cityid=95",
                                "file": "测试酒店/2026-09-03_2026-09-04.json",
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (input_dir / "room.json").write_text(
                json.dumps(
                    {
                        "hotel_name": "测试酒店",
                        "check_in": "2026-09-03",
                        "check_out": "2026-09-04",
                        "responses": [
                            {
                                "status": 200,
                                "method": "POST",
                                "url": "https://m.ctrip.com/restapi/soa2/33278/getHotelRoomListInland",
                                "data": {"ok": True},
                            }
                        ],
                        "room_rows": [
                            {
                                "酒店名称": "测试酒店",
                                "入住日期": "2026-09-03",
                                "离店日期": "2026-09-04",
                                "房型": "大床房",
                                "价格": 100,
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            output_path = Path(temporary_directory) / "output" / "prices.xlsx"

            builder.build(input_dir, output_path)

            self.assertTrue(output_path.is_file())
            workbook = load_workbook(output_path, read_only=True)
            self.assertEqual(
                set(workbook.sheetnames),
                {"房型价格", "采集汇总", "接口概览", "说明"},
            )
            self.assertEqual(workbook["房型价格"].max_row, 2)
            workbook.close()

    def test_collection_export_calls_the_builder_in_process(self):
        scripts_root = PROJECT_ROOT / "scripts"
        import sys

        if str(scripts_root) not in sys.path:
            sys.path.insert(0, str(scripts_root))
        import ctrip_hotel_excel_builder
        import ctrip_hotel_prices

        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            output = root / "prices.xlsx"
            with patch.object(ctrip_hotel_excel_builder, "build") as build_mock:
                result = ctrip_hotel_prices.export_excel(root, output)

        self.assertTrue(result)
        build_mock.assert_called_once_with(root, output)


if __name__ == "__main__":
    unittest.main()
